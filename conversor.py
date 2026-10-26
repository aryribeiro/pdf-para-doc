"""Motor de conversão PDF -> DOCX, executado num PROCESSO SEPARADO do Streamlit.

Por que um processo à parte: uma conversão de 2 páginas A4 chega a ~430 MB no
Python + ~330 MB no Tesseract (medido em Docker, Python 3.13). Dentro do
processo do Streamlit essa memória fica retida pelo alocador e soma com a de
outros usuários até o contêiner de 1 GB do Streamlit Cloud ser derrubado. Num
processo filho, a memória volta inteira ao sistema quando ele termina, e o app
pode matá-lo (com o Tesseract junto) se passar do tempo ou da memória.

As duas funções de conversão (modo_copia_fiel / modo_texto_editavel) e a
limpeza de linhas são as MESMAS do aplicativo original: o resultado do DOCX
não muda. A única diferença é o teto de tempo por página no Tesseract, com
uma nova tentativa (backoff exponencial + jitter) só em timeout ou erro do
sistema operacional.

Uso (chamado pelo aplicativo.py): python conversor.py <fiel|editavel> <entrada.pdf> <saida.docx>
Códigos de saída: 0 ok, 3 erro de conversão, 4 Tesseract esgotou as tentativas.
"""
import io
import os
import random
import re
import sys
import tempfile
import time

# Uma thread de OpenMP por Tesseract: o resultado do OCR é o mesmo, e o
# contêiner compartilhado não é disputado por N threads por página.
os.environ.setdefault("OMP_THREAD_LIMIT", "1")

import pymupdf as fitz  # o nome "fitz" está obsoleto no PyMuPDF 2.0
from docx import Document
from docx.shared import Pt
from PIL import Image
from pypdf import PdfWriter
import pytesseract

OCR_TIMEOUT_S = 90       # teto por página para o Tesseract
OCR_TENTATIVAS = 2


class TesseractIndisponivel(RuntimeError):
    """O Tesseract estourou o tempo ou falhou no SO em todas as tentativas."""


def _com_retentativa(fn, *args, **kwargs):
    """Repete só em timeout ou erro do sistema; uma página que o Tesseract não
    lê falha igual na segunda vez, então não vale insistir por conteúdo."""
    for tentativa in range(OCR_TENTATIVAS):
        try:
            return fn(*args, timeout=OCR_TIMEOUT_S, **kwargs)
        except (RuntimeError, OSError) as e:
            # pytesseract sinaliza timeout com RuntimeError("Tesseract process timeout")
            if isinstance(e, RuntimeError) and "timeout" not in str(e).lower():
                raise
            if tentativa == OCR_TENTATIVAS - 1:
                raise TesseractIndisponivel(str(e)) from e
            # random: jitter de espera, não é uso criptográfico
            time.sleep(min(3.0, 0.5 * (2 ** tentativa)) + random.uniform(0, 0.3))  # nosec B311


# ==========================================
# MÓDULO 1: Cópia Fiel (Para Impressão)
# ==========================================
def modo_copia_fiel(pdf_file):
    """Gera PDF intermediário com OCR e aplica pdf2docx para manter 100% da fidelidade visual."""
    from pdf2docx import Converter  # pesado (OpenCV): só entra neste modo

    pdf_bytes = pdf_file.read()
    doc_original = fitz.open(stream=pdf_bytes, filetype="pdf")
    merger = PdfWriter()

    for page_num in range(len(doc_original)):
        page = doc_original[page_num]
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes()))
        del pix

        pdf_ocr_bytes = _com_retentativa(
            pytesseract.image_to_pdf_or_hocr, img, extension="pdf", lang="por"
        )
        img.close()
        merger.append(io.BytesIO(pdf_ocr_bytes))
    doc_original.close()

    with tempfile.NamedTemporaryFile(
        delete=False, suffix="_ocr.pdf"
    ) as temp_ocr_pdf:
        merger.write(temp_ocr_pdf)
        temp_ocr_pdf_path = temp_ocr_pdf.name
    merger.close()

    temp_docx_path = temp_ocr_pdf_path.replace(".pdf", ".docx")

    try:
        cv = Converter(temp_ocr_pdf_path)
        cv.convert(temp_docx_path, start=0, end=None)
        cv.close()

        with open(temp_docx_path, "rb") as f:
            return f.read()
    finally:
        if os.path.exists(temp_ocr_pdf_path):
            os.remove(temp_ocr_pdf_path)
        if os.path.exists(temp_docx_path):
            os.remove(temp_docx_path)


# ==========================================
# MÓDULO 2: Texto Editável (Limpeza de Artefatos)
# ==========================================
def limpar_linha(texto_linha):
    """Filtra lixo de navegação e remove caracteres/números parasitas que aparecem antes das frases devido a ícones."""
    if not texto_linha:
        return None

    l_lower = texto_linha.lower().strip()

    blacklist = ["firefox", "about:blank", "lofl", "1 of 1"]
    if any(termo in l_lower for termo in blacklist):
        return None

    # Remove marcadores e caracteres estranhos
    texto_limpo = re.sub(
        r"^\s*[\(\[\{]?\d+[\)\]\}]?\s+(?=[A-Za-zÀ-ÿ])", "", texto_linha
    )
    texto_limpo = re.sub(
        r"^\s*[\(\[\{]?[A-Za-z0-9]{1,2}[\)\]\}]?\s+(?=[A-Za-zÀ-ÿ])",
        "",
        texto_limpo,
    )
    texto_limpo = re.sub(r"^\s*[^a-zA-Z0-9À-ÿ]+\s*(?=[A-Za-zÀ-ÿ])", "", texto_limpo)

    return texto_limpo.strip() if texto_limpo.strip() else None


def modo_texto_editavel(pdf_file):
    """Extrai texto e coordenadas via OCR, limpa artefatos e constrói parágrafos nativos no Word."""
    doc_word = Document()
    pdf_bytes = pdf_file.read()
    pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")

    for page_num in range(len(pdf_document)):
        page = pdf_document[page_num]
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes()))
        del pix

        data = _com_retentativa(
            pytesseract.image_to_data,
            img, lang="por", output_type=pytesseract.Output.DICT
        )
        img.close()

        linhas = {}
        num_items = len(data["text"])

        for i in range(num_items):
            texto = data["text"][i].strip()
            confianca = int(data["conf"][i])

            if confianca > 30 and texto:
                block_num = data["block_num"][i]
                line_num = data["line_num"][i]
                chave_linha = (block_num, line_num)

                top = data["top"][i]
                height = data["height"][i]

                if chave_linha not in linhas:
                    linhas[chave_linha] = {
                        "palavras": [],
                        "top": top,
                        "heights": [],
                    }

                linhas[chave_linha]["palavras"].append(texto)
                linhas[chave_linha]["heights"].append(height)

        linhas_ordenadas = sorted(linhas.values(), key=lambda x: x["top"])

        for item in linhas_ordenadas:
            texto_bruto = " ".join(item["palavras"])
            texto_linha = limpar_linha(texto_bruto)

            if not texto_linha:
                continue

            altura_media_px = sum(item["heights"]) / len(item["heights"])
            tamanho_fonte_pt = max(9, min(26, int(altura_media_px / 3.8)))

            p = doc_word.add_paragraph()
            run = p.add_run(texto_linha)
            run.font.name = "Arial"
            run.font.size = Pt(tamanho_fonte_pt)

            if (
                tamanho_fonte_pt >= 14
                or texto_linha.isdigit()
                or "Protocolo" in texto_linha
            ):
                run.bold = True

            p.paragraph_format.space_after = Pt(4)

        if page_num < len(pdf_document) - 1:
            doc_word.add_page_break()
    pdf_document.close()

    docx_buffer = io.BytesIO()
    doc_word.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer.read()


MODOS = {"fiel": modo_copia_fiel, "editavel": modo_texto_editavel}


def main(argv):
    modo, entrada, saida = argv[1:4]
    with open(entrada, "rb") as f:
        buf = io.BytesIO(f.read())
    try:
        docx_bytes = MODOS[modo](buf)
    except TesseractIndisponivel as e:
        print(f"tesseract: {e}", file=sys.stderr)
        return 4
    except Exception as e:  # noqa: BLE001 — o pai só precisa do tipo e do código
        print(f"conversao: {type(e).__name__}: {e}", file=sys.stderr)
        return 3
    with open(saida, "wb") as f:
        f.write(docx_bytes)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
